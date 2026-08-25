from pathlib import Path

from docx import Document
from docx.enum.section import WD_ORIENT, WD_SECTION
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor


ROOT = Path(__file__).resolve().parents[1]
OUTPUT = ROOT / "Batuan_Voting_System_Editable_Flowcharts.docx"

PAGE_W = 11.0
PAGE_H = 8.5
MARGIN = 0.42
CONTENT_W_DXA = 14600
FONT = "Arial"
NAVY = "102C57"
BLUE = "2E74B5"
GOLD = "C99216"
INK = "172033"
MUTED = "536174"
CANVAS = "0B1424"
CANVAS_LINE = "34445C"
WHITE = "FFFFFF"
NODE_FILLS = {
    "TERMINAL": "E8F1FB",
    "INPUT": "FFF1BF",
    "DECISION": "E9EEF7",
    "PROCESS": "FFFFFF",
    "OUTPUT": "E4F5EC",
    "CONNECTOR": "F0E6FF",
    "NOTE": "E9EDF3",
}


def set_run_font(run, name=FONT, size=9, color=INK, bold=False, italic=False):
    run.font.name = name
    rpr = run._element.get_or_add_rPr()
    rfonts = rpr.rFonts
    if rfonts is None:
        rfonts = OxmlElement("w:rFonts")
        rpr.append(rfonts)
    rfonts.set(qn("w:ascii"), name)
    rfonts.set(qn("w:hAnsi"), name)
    rfonts.set(qn("w:eastAsia"), name)
    run.font.size = Pt(size)
    run.font.color.rgb = RGBColor.from_string(color)
    run.bold = bold
    run.italic = italic


def set_cell_shading(cell, fill):
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = tc_pr.find(qn("w:shd"))
    if shd is None:
        shd = OxmlElement("w:shd")
        tc_pr.append(shd)
    shd.set(qn("w:fill"), fill)
    shd.set(qn("w:val"), "clear")


def set_cell_borders(cell, color=CANVAS_LINE, size=6, val="single"):
    tc_pr = cell._tc.get_or_add_tcPr()
    borders = tc_pr.find(qn("w:tcBorders"))
    if borders is None:
        borders = OxmlElement("w:tcBorders")
        tc_pr.append(borders)
    for edge in ("top", "left", "bottom", "right", "insideH", "insideV"):
        tag = qn(f"w:{edge}")
        item = borders.find(tag)
        if item is None:
            item = OxmlElement(f"w:{edge}")
            borders.append(item)
        item.set(qn("w:val"), val)
        item.set(qn("w:sz"), str(size))
        item.set(qn("w:space"), "0")
        item.set(qn("w:color"), color)


def set_cell_margins(cell, top=70, bottom=70, start=95, end=95):
    tc_pr = cell._tc.get_or_add_tcPr()
    tc_mar = tc_pr.find(qn("w:tcMar"))
    if tc_mar is None:
        tc_mar = OxmlElement("w:tcMar")
        tc_pr.append(tc_mar)
    for edge, value in (("top", top), ("bottom", bottom), ("start", start), ("end", end)):
        node = tc_mar.find(qn(f"w:{edge}"))
        if node is None:
            node = OxmlElement(f"w:{edge}")
            tc_mar.append(node)
        node.set(qn("w:w"), str(value))
        node.set(qn("w:type"), "dxa")


def set_table_geometry(table, widths_dxa, indent=95):
    table.autofit = False
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    tbl = table._tbl
    tbl_pr = tbl.tblPr

    tbl_w = tbl_pr.find(qn("w:tblW"))
    if tbl_w is None:
        tbl_w = OxmlElement("w:tblW")
        tbl_pr.append(tbl_w)
    tbl_w.set(qn("w:w"), str(sum(widths_dxa)))
    tbl_w.set(qn("w:type"), "dxa")

    layout = tbl_pr.find(qn("w:tblLayout"))
    if layout is None:
        layout = OxmlElement("w:tblLayout")
        tbl_pr.append(layout)
    layout.set(qn("w:type"), "fixed")

    tbl_ind = tbl_pr.find(qn("w:tblInd"))
    if tbl_ind is None:
        tbl_ind = OxmlElement("w:tblInd")
        tbl_pr.append(tbl_ind)
    tbl_ind.set(qn("w:w"), str(indent))
    tbl_ind.set(qn("w:type"), "dxa")

    grid = tbl.tblGrid
    for child in list(grid):
        grid.remove(child)
    for width in widths_dxa:
        col = OxmlElement("w:gridCol")
        col.set(qn("w:w"), str(width))
        grid.append(col)

    for row in table.rows:
        for idx, cell in enumerate(row.cells):
            tc_pr = cell._tc.get_or_add_tcPr()
            tc_w = tc_pr.find(qn("w:tcW"))
            if tc_w is None:
                tc_w = OxmlElement("w:tcW")
                tc_pr.append(tc_w)
            tc_w.set(qn("w:w"), str(widths_dxa[idx]))
            tc_w.set(qn("w:type"), "dxa")
            cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
            set_cell_margins(cell)


def set_row_cant_split(row):
    tr_pr = row._tr.get_or_add_trPr()
    if tr_pr.find(qn("w:cantSplit")) is None:
        tr_pr.append(OxmlElement("w:cantSplit"))


def clear_cell(cell):
    cell.text = ""
    p = cell.paragraphs[0]
    p.paragraph_format.space_before = Pt(0)
    p.paragraph_format.space_after = Pt(0)
    p.paragraph_format.line_spacing = 1.0
    return p


def add_cell_lines(cell, lines, color=INK, size=8.5, bold=False, align=WD_ALIGN_PARAGRAPH.CENTER, fill=WHITE, border_color=CANVAS_LINE):
    set_cell_shading(cell, fill)
    set_cell_borders(cell, border_color, 7, "single")
    set_cell_margins(cell, top=85, bottom=85, start=95, end=95)
    cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
    cell.text = ""
    for idx, line in enumerate(lines):
        p = cell.paragraphs[0] if idx == 0 else cell.add_paragraph()
        p.alignment = align
        p.paragraph_format.space_before = Pt(0)
        p.paragraph_format.space_after = Pt(0)
        p.paragraph_format.line_spacing = 1.0
        if isinstance(line, tuple):
            text, line_size, line_bold, line_color, line_italic = line
        else:
            text, line_size, line_bold, line_color, line_italic = line, size, bold, color, False
        run = p.add_run(text)
        set_run_font(run, size=line_size, color=line_color, bold=line_bold, italic=line_italic)


def node(cell, kind, body, size=8.5):
    lines = [(kind, 6.6, True, GOLD if kind in ("DECISION", "INPUT") else NAVY, False)]
    if isinstance(body, (list, tuple)):
        lines += [(str(item), size, False, INK, False) for item in body]
    else:
        lines += [(line, size, False, INK, False) for line in str(body).split("\n")]
    add_cell_lines(cell, lines, fill=NODE_FILLS[kind], border_color="BFC9D8")


def arrow(cell, text="→", size=16):
    add_cell_lines(cell, [(text, size, True, GOLD, False)], fill=CANVAS, border_color=CANVAS)


def canvas_table(doc, rows, widths, default_fill=CANVAS):
    table = doc.add_table(rows=rows, cols=len(widths))
    set_table_geometry(table, widths)
    for row in table.rows:
        set_row_cant_split(row)
        for cell in row.cells:
            add_cell_lines(cell, [""], fill=default_fill, border_color=CANVAS_LINE)
    return table


def add_chain(doc, chain, font_size=8.4):
    widths = []
    for idx in range(len(chain)):
        widths.append(2250 if idx % 2 == 0 else 480)
    table = canvas_table(doc, 1, widths)
    for idx, item in enumerate(chain):
        if idx % 2 == 0:
            kind, body = item
            node(table.cell(0, idx), kind, body, font_size)
        else:
            arrow(table.cell(0, idx), item)
    doc.add_paragraph().paragraph_format.space_after = Pt(1)
    return table


def add_branch_table(doc, rows, font_size=7.7):
    widths = [3200, 3850, 3850, 3700]
    table = canvas_table(doc, len(rows) + 1, widths)
    headers = ["FLOW NODE / DECISION", "YES (Y) PATH", "NO (N) PATH / NEXT", "SYSTEM DATA / GUARD"]
    for idx, header in enumerate(headers):
        add_cell_lines(table.cell(0, idx), [(header, 6.7, True, WHITE, False)], fill=NAVY, border_color=BLUE)
    for ridx, row in enumerate(rows, start=1):
        set_row_cant_split(table.rows[ridx])
        for cidx, cell_data in enumerate(row):
            cell = table.cell(ridx, cidx)
            if isinstance(cell_data, dict) and "kind" in cell_data:
                label = cell_data.get("label", "")
                body = cell_data.get("body", "")
                lines = [(label, 6.2, True, GOLD, False)] if label else []
                body_lines = body if isinstance(body, (list, tuple)) else str(body).split("\n")
                lines.extend([(str(v), font_size, False, INK, False) for v in body_lines])
                add_cell_lines(cell, lines or [""], fill=NODE_FILLS.get(cell_data["kind"], WHITE), border_color="BFC9D8")
            elif isinstance(cell_data, str):
                add_cell_lines(cell, [(cell_data, font_size, False, WHITE if cidx else INK, False)], fill=CANVAS if cidx else NODE_FILLS["NOTE"], border_color=CANVAS_LINE)
            else:
                add_cell_lines(cell, [""], fill=CANVAS, border_color=CANVAS_LINE)
    doc.add_paragraph().paragraph_format.space_after = Pt(1)
    return table


def add_note_box(doc, title, text, fill="F2F5F9"):
    table = doc.add_table(rows=1, cols=1)
    set_table_geometry(table, [CONTENT_W_DXA])
    cell = table.cell(0, 0)
    lines = [(title, 7.4, True, NAVY, False), (text, 7.6, False, INK, False)]
    add_cell_lines(cell, lines, fill=fill, border_color="C7D2E1", align=WD_ALIGN_PARAGRAPH.LEFT)
    doc.add_paragraph().paragraph_format.space_after = Pt(1)


def add_legend(doc):
    table = doc.add_table(rows=2, cols=6)
    set_table_geometry(table, [2430] * 6)
    labels = [
        ("TERMINAL", "START / END"),
        ("INPUT", "Menu / form"),
        ("DECISION", "Condition"),
        ("PROCESS", "Save / update"),
        ("OUTPUT", "Display"),
        ("CONNECTOR", "Off-page"),
    ]
    for idx, (kind, label) in enumerate(labels):
        add_cell_lines(table.cell(0, idx), [(kind, 6.3, True, NAVY, False)], fill=NODE_FILLS[kind], border_color="BFC9D8")
        add_cell_lines(table.cell(1, idx), [(label, 7, False, INK, False)], fill=NODE_FILLS[kind], border_color="BFC9D8")
    doc.add_paragraph().paragraph_format.space_after = Pt(2)


def set_page(section):
    section.orientation = WD_ORIENT.LANDSCAPE
    section.page_width = Inches(PAGE_W)
    section.page_height = Inches(PAGE_H)
    section.top_margin = Inches(MARGIN)
    section.bottom_margin = Inches(MARGIN)
    section.left_margin = Inches(MARGIN)
    section.right_margin = Inches(MARGIN)
    section.header_distance = Inches(0.2)
    section.footer_distance = Inches(0.2)


def set_header_footer(section, figure_label="BNHS SSLG | Editable Flowcharts"):
    header = section.header
    hp = header.paragraphs[0]
    hp.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    hp.paragraph_format.space_after = Pt(0)
    hp.text = ""
    r = hp.add_run(figure_label)
    set_run_font(r, size=7.2, color=MUTED, bold=True)
    footer = section.footer
    fp = footer.paragraphs[0]
    fp.alignment = WD_ALIGN_PARAGRAPH.CENTER
    fp.paragraph_format.space_before = Pt(0)
    fp.text = ""
    r = fp.add_run("Batuan National High School • SSLG Voting & Management System")
    set_run_font(r, size=7.2, color=MUTED)


def add_title(doc, title, subtitle=None):
    p = doc.add_paragraph(style="Heading 1")
    p.paragraph_format.space_before = Pt(0)
    p.paragraph_format.space_after = Pt(1)
    r = p.add_run(title)
    set_run_font(r, size=16, color=NAVY, bold=True)
    if subtitle:
        p2 = doc.add_paragraph()
        p2.paragraph_format.space_before = Pt(0)
        p2.paragraph_format.space_after = Pt(4)
        r2 = p2.add_run(subtitle)
        set_run_font(r2, size=8.5, color=MUTED, italic=True)


def add_figure_heading(doc, fig_no, title, subtitle):
    add_title(doc, f"Figure {fig_no}. {title}", subtitle)


def add_section_for_figure(doc, figure_label):
    section = doc.add_section(WD_SECTION.NEW_PAGE)
    set_page(section)
    set_header_footer(section, figure_label)
    return section


def add_cover(doc):
    section = doc.sections[0]
    set_page(section)
    set_header_footer(section, "BNHS SSLG | System Flowchart Specification")
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_before = Pt(22)
    p.paragraph_format.space_after = Pt(1)
    r = p.add_run("BATUAN NATIONAL HIGH SCHOOL")
    set_run_font(r, size=14, color=BLUE, bold=True)
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_before = Pt(0)
    p.paragraph_format.space_after = Pt(2)
    r = p.add_run("SUPREME SECONDARY LEARNER GOVERNMENT (SSLG)")
    set_run_font(r, size=20, color=NAVY, bold=True)
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_before = Pt(0)
    p.paragraph_format.space_after = Pt(6)
    r = p.add_run("Automated Voting & Management System")
    set_run_font(r, size=17, color=NAVY, bold=True)
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_before = Pt(0)
    p.paragraph_format.space_after = Pt(14)
    r = p.add_run("Editable Word-native Flowchart Set | Version 1.0")
    set_run_font(r, size=10, color=MUTED, italic=True)

    meta = doc.add_table(rows=4, cols=2)
    set_table_geometry(meta, [2850, 11750])
    items = [
        ("System basis", "Active React/Vite routes, Supabase queries/RPCs, and PostgreSQL schema in the batuan-voting workspace."),
        ("Reference basis", "CargoExpressFlowchart.docx and the CargoExpressPH source tree; Figure 2–8 decomposition retained as the organizing pattern."),
        ("Editable format", "All flowchart nodes are Word table cells with editable text and formatting. No diagram screenshot is embedded."),
        ("Primary roles", "Public visitor, student voter, and administrator / election manager."),
    ]
    for i, (k, v) in enumerate(items):
        add_cell_lines(meta.cell(i, 0), [(k.upper(), 7.5, True, NAVY, False)], fill="E8EEF5", border_color="C7D2E1", align=WD_ALIGN_PARAGRAPH.LEFT)
        add_cell_lines(meta.cell(i, 1), [(v, 8, False, INK, False)], fill="FFFFFF", border_color="C7D2E1", align=WD_ALIGN_PARAGRAPH.LEFT)
    doc.add_paragraph().paragraph_format.space_after = Pt(6)

    add_title(doc, "Flowchart symbol key", "The diagram pages use a dark canvas with light, editable node cells to echo the supplied reference while keeping content customizable in Word.")
    add_legend(doc)
    add_note_box(doc, "Customization note", "To customize: click inside any node cell, replace the text, then adjust the cell fill/border if needed. Use the Y/N columns as branch labels; duplicate a row when adding a new decision. The structure is intentionally table-based so it remains editable in Word, LibreOffice, and Google Docs import.", fill="FFF7DE")
    add_note_box(doc, "Current-system accuracy note", "The active client routes data through src/api/client.js to Supabase public reads and app_* PostgreSQL RPCs. The legacy README wording about Express/MySQL is not used as the source of truth for these flowcharts.", fill="EDF5FF")


def build():
    doc = Document()
    styles = doc.styles
    normal = styles["Normal"]
    normal.font.name = "Calibri"
    normal.font.size = Pt(9.5)
    normal.paragraph_format.space_after = Pt(3)
    normal.paragraph_format.line_spacing = 1.1
    for style_name, size, color in (("Heading 1", 16, NAVY), ("Heading 2", 12, BLUE), ("Heading 3", 10.5, NAVY)):
        style = styles[style_name]
        style.font.name = FONT
        style.font.size = Pt(size)
        style.font.color.rgb = RGBColor.from_string(color)
        style.font.bold = True
        style.paragraph_format.space_before = Pt(3)
        style.paragraph_format.space_after = Pt(3)
    add_cover(doc)

    # Figure 2
    add_section_for_figure(doc, "BNHS SSLG | Figure 2 | Public & Authentication")
    add_figure_heading(doc, 2, "Public Landing Page & Authentication", "Entry routes: /, /auth, /change-password | Branch labels Y = yes, N = no")
    add_chain(doc, [("TERMINAL", "START"), "→", ("OUTPUT", "Load public landing page\n/"), "→", ("INPUT", "Choices:\nA. Home / Overview\nB. Candidates\nC. Results\nD. Sign In")], 8.3)
    rows = [
        ({"kind": "DECISION", "label": "A", "body": "if choice == Home / Overview"}, {"kind": "OUTPUT", "label": "Y", "body": "Display election status, date, turnout stats, leading candidates, and quick actions"}, {"kind": "DECISION", "label": "N → next", "body": "Check choice B"}, {"kind": "PROCESS", "label": "GET", "body": "/election-settings\n/stats\n/candidates\n/positions"}),
        ({"kind": "DECISION", "label": "B", "body": "if choice == Candidates"}, {"kind": "OUTPUT", "label": "Y", "body": "Display active candidate directory; search by name / party / section and filter by position"}, {"kind": "DECISION", "label": "N → next", "body": "Check choice C"}, {"kind": "PROCESS", "label": "GET", "body": "/candidates\n/positions"}),
        ({"kind": "DECISION", "label": "C", "body": "if choice == Results"}, {"kind": "OUTPUT", "label": "Y", "body": "Display Live Results or Past Elections with position, grade, and section filters"}, {"kind": "DECISION", "label": "N → next", "body": "Check choice D"}, {"kind": "PROCESS", "label": "GET", "body": "/votes/counts\n/election-history"}),
        ({"kind": "DECISION", "label": "D", "body": "if choice == Sign In"}, {"kind": "INPUT", "label": "Y", "body": "Enter LRN / admin username + password"}, {"kind": "INPUT", "label": "N → loop", "body": "Return to public choices"}, {"kind": "PROCESS", "label": "POST", "body": "/auth/login → app_login"}),
        ({"kind": "DECISION", "label": "1", "body": "Credentials valid?"}, {"kind": "DECISION", "label": "Y", "body": "Check role"}, {"kind": "OUTPUT", "label": "N", "body": "Show login error toast; retry credentials"}, {"kind": "PROCESS", "label": "AUTH", "body": "Store auth_token; fetch /auth/me"}),
        ({"kind": "DECISION", "label": "2", "body": "isAdmin == true?"}, {"kind": "CONNECTOR", "label": "Y → 1", "body": "Admin Dashboard /admin"}, {"kind": "DECISION", "label": "N", "body": "Check must_change_password"}, {"kind": "PROCESS", "label": "ROLE", "body": "user_roles = admin or voter"}),
        ({"kind": "DECISION", "label": "3", "body": "must_change_password?"}, {"kind": "INPUT", "label": "Y", "body": "Change password (minimum 6 chars)"}, {"kind": "CONNECTOR", "label": "N → 2", "body": "Student Voter Dashboard /"}, {"kind": "PROCESS", "label": "SAVE", "body": "app_change_password\nclear flag; refresh session"}),
        ({"kind": "OUTPUT", "label": "SUCCESS", "body": "Password saved → Student Voter Dashboard"}, {"kind": "CONNECTOR", "label": "→ 2", "body": "Voter portal"}, {"kind": "NOTE", "label": "", "body": "Public routes remain accessible without a session"}, {"kind": "PROCESS", "label": "SECURITY", "body": "Invalid token → clear local auth_token"}),
    ]
    add_branch_table(doc, rows, 7.1)
    add_note_box(doc, "Implementation anchors", "src/App.jsx routes /, /auth, /change-password; src/contexts/AuthContext.jsx handles signIn, fetchMe, role, must_change_password, and signOut.", fill="F2F5F9")

    # Figure 3
    add_section_for_figure(doc, "BNHS SSLG | Figure 3 | Student Voter Portal")
    add_figure_heading(doc, 3, "Student Voter Dashboard & Ballot Casting", "Student route: / | Voting route: /vote | Candidate route: /candidates | Results route: /results")
    add_chain(doc, [("CONNECTOR", "2"), "→", ("OUTPUT", "Student Voter Dashboard\n/"), "→", ("INPUT", "Choices:\n1. Dashboard\n2. Candidates\n3. Cast Vote\n4. Results\n5. Account"), "→", ("OUTPUT", "Display voter status\nNot voted / Voted")], 7.9)
    rows = [
        ({"kind": "DECISION", "label": "1", "body": "Dashboard"}, {"kind": "OUTPUT", "label": "Y", "body": "Show election status, date, turnout, voter status, and quick actions"}, {"kind": "DECISION", "label": "N → next", "body": "Check choice 2"}, {"kind": "PROCESS", "label": "GET", "body": "/election-settings\n/stats\n/auth/me"}),
        ({"kind": "DECISION", "label": "2", "body": "Candidates"}, {"kind": "OUTPUT", "label": "Y", "body": "Browse candidate cards; filter by position; view party and motto"}, {"kind": "DECISION", "label": "N → next", "body": "Check choice 3"}, {"kind": "PROCESS", "label": "GET", "body": "/candidates\n/positions"}),
        ({"kind": "DECISION", "label": "3", "body": "Cast Vote"}, {"kind": "DECISION", "label": "Y", "body": "Check election status"}, {"kind": "DECISION", "label": "N → next", "body": "Check choice 4"}, {"kind": "PROCESS", "label": "GUARD", "body": "User required; admins cannot vote"}),
        ({"kind": "DECISION", "label": "4", "body": "status == ongoing?"}, {"kind": "DECISION", "label": "Y", "body": "Check profile.has_voted"}, {"kind": "OUTPUT", "label": "N", "body": "Show Upcoming / Voting Has Ended; link to candidates or results"}, {"kind": "PROCESS", "label": "STATE", "body": "upcoming → ongoing → completed"}),
        ({"kind": "DECISION", "label": "5", "body": "has_voted?"}, {"kind": "OUTPUT", "label": "Y", "body": "Display You've Already Voted; link to Results"}, {"kind": "INPUT", "label": "N", "body": "Load ballot positions and candidates"}, {"kind": "PROCESS", "label": "FILTER", "body": "Grade Representative positions match voter grade"}),
        ({"kind": "INPUT", "label": "BALLOT", "body": "Select candidates per position"}, {"kind": "PROCESS", "label": "Y", "body": "Selection count <= position.max_votes"}, {"kind": "NOTE", "label": "N", "body": "Toast: max selection reached; deselect first"}, {"kind": "DATA", "label": "RULE", "body": "P.I.O. and Peace Officer allow up to 2"}),
        ({"kind": "PROCESS", "label": "SUBMIT", "body": "Submit Final Vote"}, {"kind": "OUTPUT", "label": "SUCCESS", "body": "Vote Submitted; refresh profile and results"}, {"kind": "OUTPUT", "label": "ERROR", "body": "Show vote failed toast; remain on ballot"}, {"kind": "PROCESS", "label": "RPC", "body": "app_submit_votes → insert votes; set has_voted = true"}),
        ({"kind": "DECISION", "label": "5.1 / 5.2", "body": "Account: Change Password or Sign Out"}, {"kind": "PROCESS", "label": "Y", "body": "Save password or clear local auth_token"}, {"kind": "INPUT", "label": "N", "body": "Return to account choices"}, {"kind": "PROCESS", "label": "END", "body": "Sign out clears user/profile/admin state"}),
    ]
    add_branch_table(doc, rows, 7.0)
    add_note_box(doc, "Implementation anchors", "src/pages/VotePage.jsx enforces client guards and selection limits; server/schema.sql app_submit_votes enforces authentication, election state, grade representative matching, max_votes, and duplicate-vote protection.", fill="F2F5F9")

    # Figure 4
    add_section_for_figure(doc, "BNHS SSLG | Figure 4 | Admin Dashboard")
    add_figure_heading(doc, 4, "Administrator Dashboard & Navigation Hub", "Admin panel tabs 1–5; global navigation also exposes Results and Sign Out")
    add_chain(doc, [("CONNECTOR", "1"), "→", ("OUTPUT", "Admin Dashboard\n/admin"), "→", ("INPUT", "Choices:\n1. Overview\n2. Voters\n3. Candidates\n4. Archive\n5. Settings\n6. Results (global nav)\n7. Sign Out")], 8.0)
    rows = [
        ({"kind": "DECISION", "label": "1", "body": "Overview"}, {"kind": "OUTPUT", "label": "Y", "body": "Display registered voters, voted count, turnout, and candidate count"}, {"kind": "DECISION", "label": "N → next", "body": "Check choice 2"}, {"kind": "PROCESS", "label": "GET", "body": "/stats\n/voters\n/candidates"}),
        ({"kind": "DECISION", "label": "2", "body": "Voters"}, {"kind": "CONNECTOR", "label": "Y → 2", "body": "Voter Management"}, {"kind": "DECISION", "label": "N → next", "body": "Check choice 3"}, {"kind": "PROCESS", "label": "AUTH", "body": "require_admin RPC guard"}),
        ({"kind": "DECISION", "label": "3", "body": "Candidates"}, {"kind": "CONNECTOR", "label": "Y → 3", "body": "Candidate Management"}, {"kind": "DECISION", "label": "N → next", "body": "Check choice 4"}, {"kind": "PROCESS", "label": "AUTH", "body": "Admin-only mutations"}),
        ({"kind": "DECISION", "label": "4", "body": "Archive"}, {"kind": "CONNECTOR", "label": "Y → 4", "body": "Archive / Recovery"}, {"kind": "DECISION", "label": "N → next", "body": "Check choice 5"}, {"kind": "PROCESS", "label": "SOFT", "body": "Archived flag hides item from active lists"}),
        ({"kind": "DECISION", "label": "5", "body": "Settings"}, {"kind": "CONNECTOR", "label": "Y → 5", "body": "Election Settings & History"}, {"kind": "DECISION", "label": "N → next", "body": "Check choice 6"}, {"kind": "PROCESS", "label": "UPDATE", "body": "Status, schedule, info, archive history"}),
        ({"kind": "DECISION", "label": "6", "body": "Results from global nav"}, {"kind": "CONNECTOR", "label": "Y → 6", "body": "Live / Past Results"}, {"kind": "DECISION", "label": "N → next", "body": "Check choice 7"}, {"kind": "NOTE", "label": "NAV", "body": "Results is not an Admin tab; it is a shared /results route"}),
        ({"kind": "DECISION", "label": "7", "body": "Sign Out"}, {"kind": "PROCESS", "label": "Y", "body": "Open confirmation; clear auth_token and session state"}, {"kind": "INPUT", "label": "N", "body": "Return to dashboard"}, {"kind": "TERMINAL", "label": "END", "body": "Public landing page"}),
    ]
    add_branch_table(doc, rows, 7.1)
    add_note_box(doc, "Implementation anchors", "src/pages/Admin.jsx defines the Overview, Voters, Candidates, Archive, and Settings tabs. src/components/Layout.jsx provides the shared Candidates, Results, Admin, and Sign In/Sign Out navigation.", fill="F2F5F9")

    # Figure 5A
    add_section_for_figure(doc, "BNHS SSLG | Figure 5 | Voters Management")
    add_figure_heading(doc, 5, "Voters and Candidates Management", "Connector 2 — Voter Management | All nodes are editable Word cells")
    add_chain(doc, [("CONNECTOR", "2"), "→", ("OUTPUT", "Voter Management"), "→", ("INPUT", "Search by LRN, name, or section"), "→", ("OUTPUT", "Active voters table\nstatus: Voted / Not voted / New")], 7.8)
    rows = [
        ({"kind": "INPUT", "label": "2.1", "body": "Bulk CSV Upload"}, {"kind": "PROCESS", "label": "Y", "body": "Parse and preview rows: lrn, full_name, grade_level, section"}, {"kind": "OUTPUT", "label": "N / invalid", "body": "Show row errors; keep preview for correction"}, {"kind": "PROCESS", "label": "RPC", "body": "app_bulk_upload_voters\ninsert / skip duplicate / return errors"}),
        ({"kind": "INPUT", "label": "2.2", "body": "Add New Voter"}, {"kind": "PROCESS", "label": "SAVE", "body": "Create user + profile + voter role"}, {"kind": "OUTPUT", "label": "ERROR", "body": "LRN must be exactly 12 digits; duplicate rejected"}, {"kind": "PROCESS", "label": "DEFAULT", "body": "Password = LRN; must_change_password = true"}),
        ({"kind": "INPUT", "label": "2.3", "body": "Edit voter"}, {"kind": "PROCESS", "label": "SAVE", "body": "Update LRN, name, grade, section"}, {"kind": "OUTPUT", "label": "ERROR", "body": "Reject missing fields or LRN conflict"}, {"kind": "PROCESS", "label": "RPC", "body": "app_update_voter"}),
        ({"kind": "INPUT", "label": "2.4", "body": "Reset password"}, {"kind": "PROCESS", "label": "SAVE", "body": "Reset password to LRN; force change on next login"}, {"kind": "OUTPUT", "label": "ERROR", "body": "Voter not found"}, {"kind": "PROCESS", "label": "RPC", "body": "app_reset_voter_password"}),
        ({"kind": "INPUT", "label": "2.5", "body": "Archive voter"}, {"kind": "PROCESS", "label": "SAVE", "body": "Set profiles.archived = true; remove from active list"}, {"kind": "OUTPUT", "label": "CANCEL", "body": "Keep voter active"}, {"kind": "CONNECTOR", "label": "→ 4", "body": "Available in Archive for restore / permanent delete"}),
        ({"kind": "INPUT", "label": "2.6", "body": "Reset all voting statuses"}, {"kind": "PROCESS", "label": "CONFIRM", "body": "Set has_voted = false and delete live votes"}, {"kind": "OUTPUT", "label": "CANCEL", "body": "No changes"}, {"kind": "PROCESS", "label": "NEW ELECTION", "body": "Admin confirmation required; invalidates current live tallies"}),
    ]
    add_branch_table(doc, rows, 7.1)
    add_note_box(doc, "Data guard", "All voter mutations call require_admin through app_* RPCs. The UI previews bulk uploads and reports inserted, skipped duplicate, and error counts.", fill="FFF7DE")

    # Figure 5B
    add_section_for_figure(doc, "BNHS SSLG | Figure 5 | Candidates Management")
    add_figure_heading(doc, 5, "Voters and Candidates Management", "Connector 3 — Candidate Management | Figure 5 continues on this page")
    add_chain(doc, [("CONNECTOR", "3"), "→", ("OUTPUT", "Candidate Management"), "→", ("INPUT", "Search by name, position, party, or section"), "→", ("OUTPUT", "Active candidates table")], 7.8)
    rows = [
        ({"kind": "INPUT", "label": "3.1", "body": "Bulk CSV Upload"}, {"kind": "PROCESS", "label": "Y", "body": "Parse and preview name, position, grade, section, party, motto"}, {"kind": "OUTPUT", "label": "N / invalid", "body": "Show missing fields or invalid position; keep preview"}, {"kind": "PROCESS", "label": "RPC", "body": "app_bulk_upload_candidates\ninsert / skip duplicate / return errors"}),
        ({"kind": "INPUT", "label": "3.2", "body": "Add Candidate"}, {"kind": "PROCESS", "label": "SAVE", "body": "Validate form and create candidate record"}, {"kind": "OUTPUT", "label": "ERROR", "body": "Require name, position, grade, section, party list"}, {"kind": "PROCESS", "label": "PHOTO", "body": "Optional image upload to candidate-photos storage"}),
        ({"kind": "INPUT", "label": "3.3", "body": "Edit candidate"}, {"kind": "PROCESS", "label": "SAVE", "body": "Update profile fields and optionally replace photo"}, {"kind": "OUTPUT", "label": "ERROR", "body": "Show validation or RPC error"}, {"kind": "PROCESS", "label": "RPC", "body": "app_update_candidate"}),
        ({"kind": "INPUT", "label": "3.4", "body": "Archive candidate"}, {"kind": "PROCESS", "label": "CONFIRM", "body": "Set archived = true; remove from active ballot"}, {"kind": "OUTPUT", "label": "CANCEL", "body": "Keep candidate active"}, {"kind": "CONNECTOR", "label": "→ 4", "body": "Archived candidates can be restored or permanently deleted"}),
        ({"kind": "OUTPUT", "label": "RESULT", "body": "Refresh active candidates list"}, {"kind": "PROCESS", "label": "CACHE", "body": "Invalidate candidates queries"}, {"kind": "NOTE", "label": "", "body": "Candidate position comes from positions table"}, {"kind": "PROCESS", "label": "DATA", "body": "candidates.position_id → positions.id"}),
    ]
    add_branch_table(doc, rows, 7.2)
    add_note_box(doc, "Data guard", "The active API client uses app_add_candidate/app_update_candidate RPCs and Supabase Storage for candidate photos. Candidate rows are filtered by archived = false for public and ballot views.", fill="FFF7DE")

    # Figure 6
    add_section_for_figure(doc, "BNHS SSLG | Figure 6 | Archive & Recovery")
    add_figure_heading(doc, 6, "Archive and Recovery Management", "Connector 4 — soft-delete safety flow for voters and candidates")
    add_chain(doc, [("CONNECTOR", "4"), "→", ("OUTPUT", "Archive Management Panel"), "→", ("INPUT", "Sub-tabs:\n4.1 Archived Voters\n4.2 Archived Candidates")], 8.1)
    rows = [
        ({"kind": "DECISION", "label": "4.1", "body": "Archived Voters"}, {"kind": "OUTPUT", "label": "Y", "body": "Display archived voters table; search by LRN, name, section"}, {"kind": "DECISION", "label": "N → 4.2", "body": "Open archived candidates"}, {"kind": "PROCESS", "label": "GET", "body": "app_list_archived_voters"}),
        ({"kind": "DECISION", "label": "V-A", "body": "Select voter action"}, {"kind": "PROCESS", "label": "Restore", "body": "Set archived = false; return to active voter list"}, {"kind": "PROCESS", "label": "Permanent delete", "body": "Confirm then hard delete user and cascaded records"}, {"kind": "PROCESS", "label": "RPC", "body": "app_restore_voter\napp_permanent_delete_voter"}),
        ({"kind": "DECISION", "label": "4.2", "body": "Archived Candidates"}, {"kind": "OUTPUT", "label": "Y", "body": "Display archived candidates table; search by name, position, party"}, {"kind": "DECISION", "label": "N → loop", "body": "Return to archive choices"}, {"kind": "PROCESS", "label": "GET", "body": "app_list_archived_candidates"}),
        ({"kind": "DECISION", "label": "C-A", "body": "Select candidate action"}, {"kind": "PROCESS", "label": "Restore", "body": "Set archived = false; return to active candidate list"}, {"kind": "PROCESS", "label": "Permanent delete", "body": "Confirm then remove candidate record"}, {"kind": "PROCESS", "label": "RPC", "body": "app_restore_candidate\napp_permanent_delete_candidate"}),
        ({"kind": "OUTPUT", "label": "DONE", "body": "Toast success; invalidate archive and active-list queries"}, {"kind": "CONNECTOR", "label": "→ 4", "body": "Remain in archive panel"}, {"kind": "NOTE", "label": "", "body": "Cancel leaves item unchanged"}, {"kind": "SECURITY", "label": "ADMIN", "body": "All actions require admin token"}),
    ]
    add_branch_table(doc, rows, 7.1)
    add_note_box(doc, "Safety distinction", "Archive is reversible soft delete. Permanent delete is destructive and confirmation-gated; database foreign keys cascade related records where defined.", fill="FFF0F0")

    # Figure 7
    add_section_for_figure(doc, "BNHS SSLG | Figure 7 | Settings & History")
    add_figure_heading(doc, 7, "Election Settings, Schedule, and History", "Connector 5 — election lifecycle controls and historical result snapshots")
    add_chain(doc, [("CONNECTOR", "5"), "→", ("OUTPUT", "Election Settings Panel"), "→", ("INPUT", "Choices:\n5.1 Status Control\n5.2 Schedule\n5.3 Election Info\n5.4 Archive Results\n5.5 Delete Archived Year")], 7.8)
    rows = [
        ({"kind": "DECISION", "label": "5.1", "body": "Status Control"}, {"kind": "PROCESS", "label": "SAVE", "body": "Set Upcoming / Ongoing / Completed"}, {"kind": "OUTPUT", "label": "ERROR", "body": "Show schedule or transition error"}, {"kind": "PROCESS", "label": "START", "body": "Starting Ongoing clears live votes and resets voter statuses"}),
        ({"kind": "INPUT", "label": "5.2", "body": "Schedule Configuration"}, {"kind": "PROCESS", "label": "SAVE", "body": "Save election date, voting start, end, auto-end toggle"}, {"kind": "OUTPUT", "label": "ERROR", "body": "Keep current settings; show error toast"}, {"kind": "PROCESS", "label": "AUTO", "body": "app_auto_manage_elections moves Upcoming → Ongoing → Completed"}),
        ({"kind": "INPUT", "label": "5.3", "body": "Election Info"}, {"kind": "PROCESS", "label": "SAVE", "body": "Save election name, school year, school name/location"}, {"kind": "OUTPUT", "label": "ERROR", "body": "Keep form open"}, {"kind": "PROCESS", "label": "UI", "body": "Updates page titles, hero, footer, and results letterhead"}),
        ({"kind": "INPUT", "label": "5.4", "body": "Archive current results"}, {"kind": "PROCESS", "label": "CONFIRM", "body": "Snapshot candidate ranks, winners, vote counts, and voter groups"}, {"kind": "OUTPUT", "label": "ERROR", "body": "Show failed-to-archive toast"}, {"kind": "PROCESS", "label": "RPC", "body": "app_archive_election_results"}),
        ({"kind": "INPUT", "label": "5.5", "body": "Delete archived school year"}, {"kind": "PROCESS", "label": "CONFIRM", "body": "Delete archive rows for selected school year"}, {"kind": "OUTPUT", "label": "CANCEL", "body": "Leave history unchanged"}, {"kind": "PROCESS", "label": "RPC", "body": "app_delete_election_history"}),
        ({"kind": "OUTPUT", "label": "DONE", "body": "Refresh settings and history queries; public Past Elections can read snapshots"}, {"kind": "CONNECTOR", "label": "→ 6", "body": "Results consumes archived data"}, {"kind": "NOTE", "label": "", "body": "History is separate from current live votes"}, {"kind": "DATA", "label": "TABLES", "body": "election_settings\nelection_results_archive\nelection_voter_groups_archive"}),
    ]
    add_branch_table(doc, rows, 7.0)
    add_note_box(doc, "Implementation anchors", "src/pages/Admin.jsx owns status, schedule, election info, archive, and delete-history mutations. server/migration-election-history.sql defines archive tables and history RPCs.", fill="F2F5F9")

    # Figure 8
    add_section_for_figure(doc, "BNHS SSLG | Figure 8 | Results & Analytics")
    add_figure_heading(doc, 8, "Election Results and Analytics", "Connector 6 — shared /results route for live tally, history, filtering, and print")
    add_chain(doc, [("CONNECTOR", "6"), "→", ("OUTPUT", "Election Results\n/results"), "→", ("INPUT", "Tabs:\nLive Results\nPast Elections"), "→", ("OUTPUT", "Position groups with ranks, vote counts, percentages, and winners")], 7.8)
    rows = [
        ({"kind": "DECISION", "label": "6.1", "body": "Live Results"}, {"kind": "OUTPUT", "label": "Y", "body": "Display current tallies grouped by position; highlight leading candidates"}, {"kind": "DECISION", "label": "N → 6.4", "body": "Open Past Elections"}, {"kind": "PROCESS", "label": "REALTIME", "body": "Subscribe to votes INSERT on channel live-votes; invalidate vote-count query"}),
        ({"kind": "INPUT", "label": "6.2", "body": "Live filters"}, {"kind": "PROCESS", "label": "FILTER", "body": "Select position, voter grade, and section"}, {"kind": "OUTPUT", "label": "CLEAR", "body": "Show all positions / grades / sections"}, {"kind": "PROCESS", "label": "GET", "body": "app_get_filtered_vote_counts + app_get_voter_groups"}),
        ({"kind": "OUTPUT", "label": "6.3", "body": "Winners summary"}, {"kind": "PROCESS", "label": "COMPUTE", "body": "Rank candidates per position; display winner badge(s)"}, {"kind": "NOTE", "label": "", "body": "No vote count → no winner"}, {"kind": "DATA", "label": "VIEW", "body": "vote_counts grouped by position"}),
        ({"kind": "DECISION", "label": "6.4", "body": "Past Elections"}, {"kind": "INPUT", "label": "Y", "body": "Select school year; filter position / grade / section"}, {"kind": "OUTPUT", "label": "N", "body": "Show no archived results message"}, {"kind": "PROCESS", "label": "GET", "body": "app_get_election_history\napp_get_archived_results\napp_get_archived_voter_groups"}),
        ({"kind": "DECISION", "label": "6.5", "body": "History filter supported?"}, {"kind": "OUTPUT", "label": "Y", "body": "Apply archived voter grade / section breakdown"}, {"kind": "OUTPUT", "label": "N", "body": "Disable grade / section filters; keep position filter"}, {"kind": "DATA", "label": "ARCHIVE", "body": "Historical filters use captured voter groups, not today’s roster"}),
        ({"kind": "INPUT", "label": "PRINT", "body": "Admin clicks Print Results Tally"}, {"kind": "PROCESS", "label": "Y", "body": "Open browser print dialog with seal, letterhead, tally, and signatures"}, {"kind": "OUTPUT", "label": "CANCEL", "body": "Return to results"}, {"kind": "PROCESS", "label": "BROWSER", "body": "window.print(); print CSS removes app chrome"}),
    ]
    add_branch_table(doc, rows, 7.0)
    add_note_box(doc, "Implementation anchors", "src/pages/Results.jsx owns live/history tabs, Supabase Realtime, filters, winners summary, archived results, and admin-only print. Public users can view results; only admins see the print action and edit election name.", fill="F2F5F9")

    # Final appendix
    add_section_for_figure(doc, "BNHS SSLG | Appendix | Customization & Source Map")
    add_title(doc, "Customization and source map", "Use this page as a human-editable index when revising the flowchart for a new election year or school workflow.")
    rows = [
        ("Figure 2", "Public landing and authentication", "src/App.jsx; src/pages/AuthPage.jsx; src/contexts/AuthContext.jsx; src/api/client.js"),
        ("Figure 3", "Student voter dashboard and ballot casting", "src/pages/Index.jsx; src/pages/VotePage.jsx; server/schema.sql app_submit_votes"),
        ("Figure 4", "Administrator navigation hub", "src/pages/Admin.jsx; src/components/Layout.jsx"),
        ("Figure 5", "Voters and candidates management", "src/pages/Admin.jsx; server/schema.sql voter/candidate RPCs"),
        ("Figure 6", "Archive and recovery", "src/pages/Admin.jsx; server/schema.sql app_archive/app_restore/app_permanent_delete RPCs"),
        ("Figure 7", "Settings, schedule, and history", "src/pages/Admin.jsx; server/migration-election-history.sql"),
        ("Figure 8", "Live and historical results", "src/pages/Results.jsx; server/migration-election-history.sql"),
    ]
    table = doc.add_table(rows=1, cols=3)
    set_table_geometry(table, [1500, 4200, 8900])
    for i, header in enumerate(("FIGURE", "PURPOSE", "SOURCE OF TRUTH")):
        add_cell_lines(table.cell(0, i), [(header, 7.2, True, WHITE, False)], fill=NAVY, border_color=BLUE)
    for fig, purpose, source in rows:
        row = table.add_row()
        set_table_geometry(table, [1500, 4200, 8900])
        add_cell_lines(row.cells[0], [(fig, 7.5, True, NAVY, False)], fill="E8EEF5", border_color="C7D2E1")
        add_cell_lines(row.cells[1], [(purpose, 7.4, False, INK, False)], fill="FFFFFF", border_color="C7D2E1", align=WD_ALIGN_PARAGRAPH.LEFT)
        add_cell_lines(row.cells[2], [(source, 7.2, False, INK, False)], fill="FFFFFF", border_color="C7D2E1", align=WD_ALIGN_PARAGRAPH.LEFT)

    doc.add_paragraph()
    add_note_box(doc, "Editing checklist", "1) Update election names, dates, and role labels. 2) Keep Y/N branches aligned with the condition. 3) If a route or RPC changes, update the corresponding source-map row and implementation anchor. 4) Re-open the document in Word and inspect each landscape page after edits.", fill="FFF7DE")
    add_note_box(doc, "Reference-derived visual rule", "The supplied CargoExpress flowchart used dark backgrounds and light nodes with Figure 2–8 decomposition. This document keeps that visual cue but replaces screenshot-based figures with editable tables so a human can change labels, branches, and implementation notes directly.", fill="EDF5FF")

    doc.save(OUTPUT)
    print(OUTPUT)


if __name__ == "__main__":
    build()
